from pathlib import Path

import pytest
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Table
from docx import Document
from PIL import Image, ImageDraw

from utils import document_extractor
from utils.document_extractor import extract_document
from utils.resume_analyzer import ResumeAnalyzer
from sothema_quality import prepare_analysis


def make_image(path, fmt="PNG"):
    image = Image.new("RGB", (1200, 800), "white")
    draw = ImageDraw.Draw(image)
    draw.text((80, 70), "JAKE RYAN", fill="black")
    draw.text((80, 140), "jake@example.com | Software Engineer", fill="black")
    draw.text((80, 230), "EXPERIENCE", fill="black")
    draw.text((100, 290), "Developed APIs with Python and PostgreSQL", fill="black")
    draw.text((80, 390), "EDUCATION", fill="black")
    draw.text((100, 450), "Bachelor of Computer Science", fill="black")
    image.save(path, format=fmt)


def make_pdf(path: Path):
    styles = getSampleStyleSheet()
    doc = SimpleDocTemplate(str(path), pagesize=A4)
    table = Table([[Paragraph("ALICE DUPONT<br/>alice@example.com<br/>EXPERIENCE<br/>Engineer 2021-2024", styles["BodyText"]),
                    Paragraph("EDUCATION<br/>Master Computer Science 2021<br/>SKILLS<br/>Python | SQL", styles["BodyText"])]],
                  colWidths=[250, 250])
    doc.build([table, PageBreak(), Paragraph("ALICE DUPONT", styles["Title"]),
               Paragraph("EXPERIENCE EDUCATION SKILLS", styles["BodyText"])])


def make_docx(path: Path):
    doc = Document()
    doc.add_paragraph("ALICE DUPONT")
    doc.add_paragraph("Professional profile and experience")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Skills"
    table.cell(0, 1).text = "Python | SQL"
    table.cell(1, 0).text = "Education"
    table.cell(1, 1).text = "Master Computer Science"
    doc.sections[0].header.paragraphs[0].text = "CV HEADER"
    doc.sections[0].footer.paragraphs[0].text = "CV FOOTER"
    doc.save(path)


def test_image_ocr_extracts_text_and_provenance(tmp_path, monkeypatch):
    path = tmp_path / "cv.png"
    make_image(path)
    monkeypatch.setattr(document_extractor, "_tesseract_available", lambda: True)
    monkeypatch.setattr(document_extractor, "_ocr_image_frame", lambda _image, page: [
        document_extractor.SourceLine("JAKE RYAN", page, 0, "ocr_image"),
        document_extractor.SourceLine("jake@example.com | Software Engineer", page, 0, "ocr_image"),
        document_extractor.SourceLine("EXPERIENCE", page, 0, "ocr_image"),
    ])
    result = extract_document(path.name, path.read_bytes())
    assert result.file_type == "png"
    assert "JAKE RYAN" in result.text
    assert "jake@example.com" in result.text
    assert "ocr_image" in result.methods
    assert result.ocr_pages == [1]
    assert result.as_dict()["source_lines"]


def test_image_without_tesseract_is_marked_unavailable(tmp_path, monkeypatch):
    path = tmp_path / "cv.png"
    make_image(path)
    monkeypatch.setattr(document_extractor, "_tesseract_available", lambda: False)
    result = extract_document(path.name, path.read_bytes())
    assert not result.text
    assert "ocr_unavailable" in result.warnings
    assert result.empty_pages == [1]


def test_image_extension_signature_mismatch_is_rejected(tmp_path):
    path = tmp_path / "cv.jpg"
    make_image(path, "PNG")
    with pytest.raises(ValueError, match="does not match"):
        extract_document(path.name, path.read_bytes())


def test_corrupt_image_is_rejected():
    with pytest.raises(ValueError, match="corrupt|unreadable"):
        extract_document("cv.png", b"not an image")


def test_image_dimension_safety_limit_is_enforced(monkeypatch):
    from PIL import Image
    import io

    image = Image.new("RGB", (100, 100), "white")
    handle = io.BytesIO()
    image.save(handle, format="PNG")
    monkeypatch.setattr(document_extractor, "MAX_IMAGE_PIXELS", 5000)
    with pytest.raises(ValueError, match="safety limit"):
        document_extractor.validate_image("cv.png", handle.getvalue())


def test_pdf_columns_and_provenance(tmp_path):
    path = tmp_path / "columns.pdf"
    make_pdf(path)
    result = extract_document(path.name, path.read_bytes())
    assert result.page_count == 2
    assert "ALICE DUPONT" in result.text
    assert "alice@example.com" in result.text
    assert any(method == "pdf_text_columns" for method in result.methods)
    assert result.as_dict()["source_lines"]
    assert all("page" in line and "line" in line for line in result.as_dict()["source_lines"])


def test_docx_tables_headers_and_semantic_analysis(tmp_path):
    path = tmp_path / "table.docx"
    make_docx(path)
    result = extract_document(path.name, path.read_bytes())
    assert "Python | SQL" in result.text
    assert "CV HEADER" in result.text
    assert result.tables >= 1
    analysis = prepare_analysis(result.text, ResumeAnalyzer(), extraction=result)
    assert analysis["document_type"] == "resume"
    assert analysis["sothema"]["extraction"]["tables"] >= 1
    assert 0 <= analysis["sothema"]["confidence"] <= 1


def test_invalid_document_is_not_silently_accepted():
    result = extract_document("not-a-cv.docx", b"PK\x03\x04")
    assert result.warnings
